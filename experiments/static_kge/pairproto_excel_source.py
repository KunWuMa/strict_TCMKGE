"""Read-only, explicit Excel evidence input for the current PairProto paper.

Raw experiment files are used only by the explicitly invoked Excel exporter.
Paper aggregation/plotting never silently falls back to those files. The rank
worksheets, not cached presentation metrics, are the numerical source of truth.
"""
from __future__ import annotations

from functools import lru_cache
from pathlib import Path
from types import SimpleNamespace
import copy
import hashlib
import json
import os

import numpy as np

HERE = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = HERE / 'outputs' / 'pairproto_excel_20261023' / 'PairProto_All_Seeds.xlsx'
RAW_MODE = os.environ.get('PAIRPROTO_INPUT_MODE') == 'raw-export'
RESULTS_DIR = HERE / 'results'


def workbook_path():
    return Path(os.environ.get('PAIRPROTO_RESULTS_XLSX', str(DEFAULT_WORKBOOK))).resolve()


def restore_json(rows):
    """Decode typed JSON-pointer leaves, including empty containers and nulls."""
    values = {}; children = {}
    for pointer, kind, value in rows:
        pointer = '' if pointer is None else str(pointer)
        if pointer in values: raise ValueError('Duplicate JSON pointer: '+pointer)
        values[pointer] = (kind, value)
        if pointer:
            parent, key = pointer.rsplit('/',1)
            children.setdefault(parent,[]).append(key)
    def visit(pointer):
        kind, value = values[pointer]
        if kind == 'object':
            prefix = pointer + '/'
            keys = children.get(pointer,[])
            return {key.replace('~1', '/').replace('~0', '~'): visit(prefix + key) for key in keys}
        if kind == 'array':
            return [visit(pointer + '/' + str(i)) for i in range(int(value))]
        if kind == 'null': return None
        if kind == 'boolean': return bool(value)
        if kind == 'integer': return int(value)
        if kind == 'number': return float(value)
        if kind == 'string': return '' if value is None else str(value)
        raise ValueError('Unknown metadata type: ' + str(kind))
    return visit('')


class ExcelEvidence:
    def __init__(self, path):
        from openpyxl import load_workbook  # reader only; authoring uses Artifact Tool
        self.path = Path(path)
        if not self.path.is_file():
            raise FileNotFoundError(f'Excel evidence is required: {self.path}. '
                                    'Run export_pairproto_excel.py explicitly to export completed runs.')
        self.sha256 = hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.json = {}; self.array_index = {}; self.arrays = {}; self.datasets = {}
        book = load_workbook(self.path, read_only=True, data_only=True)
        try:
            metadata = {}
            for row in book['Metadata'].iter_rows(min_row=2, values_only=True):
                if row[0] is not None: metadata.setdefault(row[0], []).append(row[1:4])
            self.json = {name: restore_json(rows) for name, rows in metadata.items()}
            if self.json.get('excel_schema.json', {}).get('schema_version') != 1:
                raise ValueError('Unsupported Excel evidence schema')
            expected=self.json['excel_schema.json'].get('protocol_sha256')
            if expected and expected != hashlib.sha256((HERE/'pairproto_paper_protocol.json').read_bytes()).hexdigest():
                raise ValueError('Workbook and paper protocol do not match')
            extension=self.json['excel_schema.json'].get('five_seed_amendment_sha256')
            if extension and extension != hashlib.sha256((HERE/'pairproto_five_seed_amendment.json').read_bytes()).hexdigest():
                raise ValueError('Workbook and five-seed amendment do not match')
            for row in book['ArrayIndex'].iter_rows(min_row=2, values_only=True):
                if row[0] is None: continue
                filename, key, sheet, col, dtype, shape, encoding = row[:7]
                self.array_index.setdefault(filename, []).append(
                    (key, sheet, int(col), dtype, json.loads(shape), encoding))
            required = {}
            for items in self.array_index.values():
                for _, sheet, col, _, _, _ in items: required.setdefault(sheet, set()).add(col)
            columns = {}
            for name, wanted in required.items():
                columns[name] = {col: [] for col in wanted}
                for row in book[name].iter_rows(min_row=2, values_only=True):
                    for col in wanted:
                        if col <= len(row) and row[col-1] is not None:
                            columns[name][col].append(row[col-1])
            for name, items in self.array_index.items():
                result = {}
                for key, sheet, col, dtype, shape, encoding in items:
                    values = np.asarray(columns[sheet][col], dtype=dtype).reshape(shape)
                    if encoding == 'rank':
                        if np.any(~np.isfinite(values)) or np.any(values < 1) or np.any(values * 2 != np.rint(values * 2)):
                            raise ValueError(f'Invalid average-tie ranks: {name}/{key}')
                        values = 1. / values
                    result[key] = values
                self.arrays[name] = result
            for row in book['Queries'].iter_rows(min_row=2, values_only=True):
                if row[0] is None: continue
                source, split, index, head, relation, target, recipe, label, count = row[:9]
                self.datasets.setdefault((source, str(split)), []).append(
                    (int(index), int(head), int(relation), int(target), int(recipe), str(label), int(count)))
        finally:
            book.close()

    def dataset(self, source, split):
        rows = sorted(self.datasets[(source, str(split))])
        if [r[0] for r in rows] != list(range(len(rows))): raise ValueError('Missing/duplicate query index')
        relation_ids = sorted(set(r[2] for r in rows))
        relations = [''] * (max(relation_ids) + 1)
        ranges = {}
        for r in rows:
            if relations[r[2]] and relations[r[2]] != r[5]: raise ValueError('Inconsistent relation label')
            relations[r[2]] = r[5]
            if r[2] in ranges and len(ranges[r[2]]) != r[6]: raise ValueError('Inconsistent range size')
            ranges[r[2]] = range(r[6])  # only cardinality is consumed by paper aggregation
        return SimpleNamespace(test=np.asarray([r[1:4] for r in rows], dtype=np.int64),
            recipe_group_ids=np.asarray([r[4] for r in rows], dtype=np.int64),
            relations=relations, range_candidates=ranges,
            audit=copy.deepcopy(self.json[f'dataset_{source}_{split}.json']['audit']))


@lru_cache(maxsize=2)
def evidence():
    return ExcelEvidence(workbook_path())


class ExcelPath:
    """Small read-only path adapter for existing, explicitly named evidence reads."""
    def __init__(self, name=''): self.name = str(name)
    def __truediv__(self, name):
        if self.name: raise ValueError('Evidence records are flat, not filesystem paths')
        return ExcelPath(name)
    def exists(self):
        store = evidence()
        return self.name in store.json or self.name in store.array_index
    def read_text(self, encoding='utf-8'):
        return json.dumps(evidence().json[self.name], ensure_ascii=False, allow_nan=False)
    def with_name(self, name): return ExcelPath(name)
    def with_suffix(self, suffix): return ExcelPath(str(Path(self.name).with_suffix(suffix)))
    def __str__(self): return str(workbook_path()) + '::' + self.name


def source_root(): return RESULTS_DIR if RAW_MODE else ExcelPath()


def read_arrays(path):
    if isinstance(path, ExcelPath):
        return {key: value.copy() for key, value in evidence().arrays[path.name].items()}
    with np.load(path) as data: return {key: data[key].copy() for key in data.files}


def load_dataset(source, split):
    if RAW_MODE:
        from pair_prototype_full_suite import load_data
        return load_data(source, split)
    return evidence().dataset(source, split)


def input_provenance():
    if RAW_MODE: return {'mode': 'raw-export', 'note': 'Explicit workbook export only'}
    store = evidence()
    return {'mode': 'excel', 'workbook': str(store.path), 'sha256': store.sha256,
            'authoritative_inputs': 'Rank worksheets, Queries, ArrayIndex and typed Metadata',
            'cached_metric_sheets': 'Human-readable outputs; independently recomputed from ranks'}
