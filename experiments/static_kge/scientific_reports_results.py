"""Excel-backed Scientific Reports follow-up and consolidated evidence."""
from pathlib import Path
import argparse,copy,hashlib,json,os,subprocess,sys
ROOT=Path(__file__).resolve().parent
BASE=ROOT/"outputs/scientific_reports_20261026"
PARENT=ROOT/"outputs/paper_budget500_20261025"
DEST=BASE/"exports"
BOOK=DEST/"PairProto_SR_All_Seeds.xlsx"
P=json.loads((BASE/"protocol.json").read_text(encoding="utf-8"))
SHA=hashlib.sha256((BASE/"protocol.json").read_bytes()).hexdigest()
METRICS=["mr","mrr","hits1","hits3","hits5","hits10"]
SEEDS=P["seeds"]
def save(path,obj):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False),encoding="utf-8")
def recpath(source,seed):
    name=P["selected"][source]["model"]
    return BASE/"results"/f"{source.lower()}_neural_{name}_seed{seed}.evaluation.json"
def export_payload():
    import numpy as np
    from openpyxl import load_workbook
    parent=PARENT/"exports/PairProto_All_Experiments_500_All_Seeds.xlsx"
    assert hashlib.sha256(parent.read_bytes()).hexdigest()==P["parent_workbook_sha256"]
    os.environ["PAIRPROTO_RESULTS_XLSX"]=str(parent)
    from pairproto_excel_source import evidence
    store=evidence()
    old=load_workbook(parent,read_only=True,data_only=True)
    rows=list(old["SeedMetrics"].values);headers=list(rows[0]);seedrows=[list(x) for x in rows[1:]]
    old.close()
    for row in seedrows:
        source,phase,key=row[1],row[3],row[4]
        if phase=="main" and key=="neural_"+P["selected"][source]["model"]:
            row[3]="main_frozen"
    tables=[];meta=[];hist=[]
    for source in P["selected"]:
        data=store.dataset(source,P["training_protocol"]["datasets"][source])
        rankrows=[[i,int(q[0]),int(q[1]),int(g)] for i,(q,g) in enumerate(zip(data.test,data.recipe_group_ids))]
        for seed in SEEDS:
            path=recpath(source,seed);record=json.loads(path.read_text(encoding="utf-8"))
            assert record["protocol_sha256"]==SHA
            cfg=P["selected"][source]
            assert record["config"]=={"lr":cfg["lr"],"heads":cfg["heads"]}
            history=record["history"];best=max(history,key=lambda h:h["dev"]["mrr"])
            assert best["epoch"]==record["selected"]["epoch"]
            assert record["stop_epoch"]<=500 and (record["stop_epoch"]==500 or record["stop_epoch"]-best["epoch"]==15)
            with np.load(record["rank_path"]) as z:
                assert np.array_equal(z["heads"],data.test[:,0])
                assert np.array_equal(z["relations"],data.test[:,1])
                ranks=np.rint(2/z["rr"])/2
                for row,value in zip(rankrows,ranks):row.append(float(value))
            short={k:record[k] for k in ["source","split","model","seed","config","selected","stop_epoch","seconds","parameters","test"]}
            meta.append([f"{source}_{seed}",json.dumps(short,ensure_ascii=False)])
            for h in history:hist.append([source,cfg["model"],seed,h["epoch"],h["loss"],h["dev"]["mrr"],h["elapsed"],int(h["epoch"]==best["epoch"])])
            oldrow=next(list(x) for x in rows[1:] if x[1]==source and x[3]=="main" and x[4]=="neural_"+cfg["model"] and x[6]==seed)
            oldrow[3]="main";oldrow[9:15]=[record["test"][m] for m in METRICS]
            oldrow[16:20]=[best["epoch"],record["stop_epoch"],record["seconds"],"New-budget development-selected follow-up"]
            seedrows.append(oldrow)
        tables.append({"name":"SR_"+source+"_Ranks","headers":["Query index","Formula ID","Relation ID","Recipe group"]+[str(s) for s in SEEDS],"rows":rankrows})
    meta.insert(0,["protocol",json.dumps(P,ensure_ascii=False)])
    leaves=[]
    def flatten(record,value,pointer=""):
        if isinstance(value,dict):
            leaves.append([record,pointer,"object",None])
            for key,item in value.items():flatten(record,item,pointer+"/"+str(key).replace("~","~0").replace("/","~1"))
        elif isinstance(value,list):
            leaves.append([record,pointer,"array",len(value)])
            for index,item in enumerate(value):flatten(record,item,pointer+"/"+str(index))
        else:
            kind="null" if value is None else "boolean" if isinstance(value,bool) else "integer" if isinstance(value,int) else "number" if isinstance(value,float) else "string"
            leaves.append([record,pointer,kind,value])
    for key,value in meta:flatten(key,json.loads(value))
    tables.insert(0,{"name":"SR_Seeds","headers":headers,"rows":seedrows})
    tables.append({"name":"SR_Metadata","headers":["Record","Pointer","Type","Value"],"rows":leaves})
    tables.append({"name":"SR_Training","headers":["Source","Model","Seed","Epoch","Loss","Dev MRR","Elapsed seconds","Selected"],"rows":hist})
    save(DEST/"followup_payload.json",{"tables":tables,"output":str(DEST/"followup_sheets.xlsx"),"parent":str(parent)})
    print("PAYLOAD",len(seedrows),"seed rows",flush=True)

def summarize_model(rrs,data):
    import numpy as np
    rows={}
    for seed,rr in rrs.items():
        rank=np.rint(2/rr)/2
        rows[str(seed)]={"mr":float(rank.mean()),"mrr":float(rr.mean()),
            **{f"hits{k}":float((rank<=k).mean()) for k in [1,3,5,10]}}
    summary={key:{"mean":float(np.mean([v[key] for v in rows.values()])),
        "sample_std":float(np.std([v[key] for v in rows.values()],ddof=1))} for key in METRICS}
    macro=[float(np.mean([rr[data.test[:,1]==r].mean() for r in np.unique(data.test[:,1])])) for rr in rrs.values()]
    relations=[]
    for relation in np.unique(data.test[:,1]):
        mask=data.test[:,1]==relation
        vals={}
        for seed,rr in rrs.items():
            rank=np.rint(2/rr[mask])/2
            vals[str(seed)]={"mr":float(rank.mean()),"mrr":float(rr[mask].mean()),
                **{f"hits{k}":float((rank<=k).mean()) for k in [1,3,5,10]}}
        relations.append({"relation":data.relations[int(relation)],"queries":int(mask.sum()),"per_seed":vals})
    return {"seeds":sorted(rrs),"runs":len(rrs),"per_seed":rows,"summary":summary,
        "relation_macro_mrr":{"mean":float(np.mean(macro)),"sample_std":float(np.std(macro,ddof=1))},
        "relation_results":relations}

def recompute():
    import numpy as np
    from openpyxl import load_workbook
    from unittest.mock import patch
    os.environ.pop("PAIRPROTO_INPUT_MODE",None)
    os.environ["PAIRPROTO_RESULTS_XLSX"]=str(BOOK)
    import pairproto_excel_source as xs
    xs.evidence.cache_clear()
    import pairproto_evidence as ev
    book=load_workbook(BOOK,read_only=True,data_only=True)
    grouped={}
    for row in book["SR_Metadata"].iter_rows(min_row=2,values_only=True):
        grouped.setdefault(row[0],[]).append(row[1:4])
    metadata={key:xs.restore_json(rows) for key,rows in grouped.items()}
    assert metadata["protocol"]==P
    newarrays={};newkeys={}
    for source in P["selected"]:
        rows=list(book["SR_"+source+"_Ranks"].iter_rows(min_row=2,values_only=True))
        assert [int(x[0]) for x in rows]==list(range(len(rows)))
        newkeys[source]=np.asarray([x[1:4] for x in rows],dtype=np.int64)
        newarrays[source]={seed:1/np.asarray([x[4+i] for x in rows],dtype=np.float64) for i,seed in enumerate(SEEDS)}
    seedrows=list(book["SR_Seeds"].iter_rows(min_row=2,values_only=True))
    book.close()
    with patch("numpy.load",side_effect=AssertionError("Scientific Reports output attempted raw rank input")):
        report=ev.build_report()
        report["frozen_main"]={}
        for source,cfg in P["selected"].items():
            split=P["training_protocol"]["datasets"][source]
            data=xs.load_dataset(source,split)
            assert np.array_equal(newkeys[source][:,:2],data.test[:,:2])
            assert np.array_equal(newkeys[source][:,2],data.recipe_group_ids)
            original,pred=ev.collect(source,split,SEEDS,mechanisms=True)
            key="neural_"+cfg["model"]
            report["frozen_main"][source]={key:copy.deepcopy(original["models"][key])}
            report["sources"][source]["main"]["models"][key]=summarize_model(newarrays[source],data)
            report.setdefault("baseline_followup",{})[source]={"key":key,
                "old":copy.deepcopy(original["models"][key]),
                "new":copy.deepcopy(report["sources"][source]["main"]["models"][key]),
                "new_minus_old":ev.bootstrap(np.mean([newarrays[source][seed]-pred[key][seed]["rr"] for seed in SEEDS],axis=0),data.recipe_group_ids,"recipe")}
            a=np.asarray([pred["pairproto"][seed]["rr"] for seed in SEEDS])
            b=np.asarray([newarrays[source][seed] for seed in SEEDS])
            delta=np.mean(a-b,axis=0)
            hit=np.mean((a>=1).astype(float)-(b>=1).astype(float),axis=0)
            report["sources"][source]["main"]["comparisons"][key]={
                "seeds":SEEDS,"mrr":ev.bootstrap(delta,data.test[:,0]),
                "hits1":ev.bootstrap(hit,data.test[:,0]),
                "recipe_mrr":ev.bootstrap(delta,data.recipe_group_ids,"recipe"),
                "recipe_hits1":ev.bootstrap(hit,data.recipe_group_ids,"recipe")}
        boundary={}
        store=xs.evidence()
        for source in P["selected"]:
            data=xs.load_dataset(source,P["training_protocol"]["datasets"][source])
            arrays={seed:store.arrays[f"pairproto_paper_{source.lower()}_boundary_seed{seed}_rr.npz"] for seed in SEEDS}
            boundary[source]={key:summarize_model({seed:a[key] for seed,a in arrays.items()},data) for key in ["mean","pairproto"]}
            delta=np.mean([a["pairproto"]-a["mean"] for a in arrays.values()],axis=0)
            boundary[source]["comparison"]=ev.bootstrap(delta,data.recipe_group_ids,"recipe")
        report["all_entity_boundary"]=boundary
    for row in seedrows:
        source,split,phase,key,seed=row[1],str(row[2]),row[3],row[4],str(row[6])
        if phase=="main_frozen":model=report["frozen_main"][source][key]
        elif phase in ["main","fixed"]:model=report["sources"][source][phase]["models"][key]
        else:model=report["sources"][source]["replications"][split]["models"][key]
        for metric,value in zip(METRICS,row[9:15]):
            assert np.isclose(model["per_seed"][seed][metric],value,rtol=1e-10,atol=1e-12),(source,key,seed,metric)
    report.update(followup_protocol=P,followup_protocol_sha256=SHA,followup_metadata=metadata,
        input_provenance={"mode":"excel","workbook":str(BOOK),"sha256":hashlib.sha256(BOOK.read_bytes()).hexdigest(),
        "authoritative_inputs":"Original rank worksheets plus SR_PRIMARY_Ranks, SR_KTKP_Ranks, Queries and typed Metadata",
        "parent_workbook_sha256":P["parent_workbook_sha256"]})
    save(DEST/"scientific_reports_summary.json",report)
    save(DEST/"excel_validation.json",{"passed":True,"seed_rows_verified":len(seedrows),"no_raw_rank_input":True,
        "workbook_sha256":report["input_provenance"]["sha256"]})
    print("EXCEL RECOMPUTATION VERIFIED",len(seedrows),flush=True)
    return report

def export_workbook():
    export_payload()
    node=Path(r"C:\Users\PC\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")
    marker=Path(r"C:\Users\PC\.codex\plugins\cache\openai-primary-runtime\spreadsheets\26.1020.11127\skills\spreadsheets\container_tools\mark_artifact_operation_started.mjs")
    payload_path=DEST/"followup_payload.json"
    payload=json.loads(payload_path.read_text(encoding="utf-8"))
    (BASE/"artifacts").mkdir(exist_ok=True)
    payload["output"]=str(BASE/"artifacts/followup_sheets.xlsx")
    save(payload_path,payload)
    if not (node.is_file() and marker.is_file()):
        from portable_scientific_reports_sheets import write
        write(payload_path)
        from merge_scientific_reports_workbook import merge
        save(DEST/"workbook_merge_validation.json",merge(payload["parent"],payload["output"],BOOK))
        return
    stamp=BASE/"artifacts/artifact_operation_started.json"
    if not stamp.exists():
        subprocess.run([str(node),str(marker),"--operation-kind","create","--expected-output-count","1","--output-format","xlsx"],
            cwd=str(ROOT),check=True,creationflags=0x08000000)
        save(stamp,{"output":str(BOOK),"operation":"create","method":"Artifact-authored supplemental sheets; lossless parent-package preservation"})
    with (BASE/"logs/artifact_workbook.log").open("w",encoding="utf-8") as log:
        subprocess.run([str(node),"--max-old-space-size=6000",str(ROOT/"build_scientific_reports_excel.mjs"),str(payload_path)],
            cwd=str(ROOT),stdout=log,stderr=subprocess.STDOUT,check=True,creationflags=0x08000000)
    from merge_scientific_reports_workbook import merge
    checked=merge(payload["parent"],payload["output"],BOOK)
    save(DEST/"workbook_merge_validation.json",checked)
    print("CONSOLIDATED WORKBOOK EXPORTED",BOOK,flush=True)

def main():
    ap=argparse.ArgumentParser();ap.add_argument("--recompute-only",action="store_true")
    ap.add_argument("--export-only",action="store_true");ap.add_argument("--statistics-only",action="store_true");a=ap.parse_args()
    if not a.recompute_only:
        for source in P["selected"]:
            for seed in SEEDS:assert recpath(source,seed).is_file(),(source,seed)
        export_workbook()
    if a.export_only:return
    recompute()
    if a.statistics_only:return
    env=os.environ.copy();env["PAIRPROTO_RESULTS_XLSX"]=str(BOOK);env.pop("PAIRPROTO_INPUT_MODE",None)
    subprocess.run([sys.executable,str(ROOT/"scientific_reports_outputs.py")],cwd=str(ROOT),env=env,
        check=True,creationflags=0x08000000)
if __name__=="__main__":main()
