"""Explicitly export completed experiment evidence into the paper's Excel input.

This is the only raw-cache -> Excel boundary. It never trains a model or chooses
hyperparameters. Spreadsheet authoring is performed by the bundled Artifact Tool.
"""
from __future__ import annotations
import os
os.environ['PAIRPROTO_INPUT_MODE'] = 'raw-export'
import argparse
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys
import numpy as np

from pairproto_evidence import PLAN, NAMES, SOURCES, collect
from pair_prototype_full_suite import load_data
from pair_memory_dev import ingredients

HERE = Path(__file__).resolve().parent
OUT = HERE / 'results'
DEFAULT_DEST = HERE / 'outputs' / 'pairproto_excel_20261023'
METRICS = ('mr', 'mrr', 'hits1', 'hits3', 'hits5', 'hits10')


def flatten(value, pointer=''):
    if isinstance(value, dict):
        yield pointer, 'object', len(value)
        for key, child in value.items():
            escaped = str(key).replace('~', '~0').replace('/', '~1')
            yield from flatten(child, pointer + '/' + escaped)
    elif isinstance(value, list):
        yield pointer, 'array', len(value)
        for i, child in enumerate(value): yield from flatten(child, pointer + '/' + str(i))
    elif value is None: yield pointer, 'null', None
    elif isinstance(value, bool): yield pointer, 'boolean', value
    elif isinstance(value, int): yield pointer, 'integer', value if abs(value) < 2**53 else str(value)
    elif isinstance(value, float):
        if not np.isfinite(value): raise ValueError('Non-finite metadata: ' + pointer)
        yield pointer, 'number', value
    else:
        text = str(value)
        if len(text) > 32000: raise ValueError('Metadata string exceeds Excel cell limit: ' + pointer)
        yield pointer, 'string', text


def identify(filename):
    source = 'KTKP' if 'ktkp' in filename else ('PRIMARY' if 'primary' in filename else 'historical')
    split_match = re.search(r'(recipe_\d+|hybrid_holdout|ktkp_(\d{8}))', filename)
    split = split_match.group(0).removeprefix('ktkp_') if split_match else 'historical'
    if split == 'historical' and source in PLAN['datasets'] and filename.startswith(('pairproto_herb_control_', 'pairproto_paper_')):
        split = str(PLAN['datasets'][source]['main'])
    seed_match = re.search(r'seed(\d+)', filename)
    return source, split, int(seed_match.group(1)) if seed_match else None


def metric_values(rr):
    rr = np.asarray(rr, dtype=float)
    return [float((1/rr).mean()), float(rr.mean())] + [float((rr >= 1/k).mean()) for k in (1,3,5,10)]


def recorded_metrics(value, pointer=''):
    if isinstance(value,dict):
        if isinstance(value.get('mrr'),(int,float)):
            yield pointer,value
        for key,child in value.items():
            yield from recorded_metrics(child,pointer+'/'+str(key))
    elif isinstance(value,list):
        for i,child in enumerate(value):yield from recorded_metrics(child,pointer+'/'+str(i))


def make_payload(dest):
    tables = []
    seed_rows = []; relation_rows = []; query_rows = []; metadata = []; source_rows = []
    seen_datasets = set()
    for source, spec in PLAN['datasets'].items():
        contexts = [('main', spec['main'], False, True), ('fixed', spec['main'], True, False)]
        contexts += [('replication', split, False, False) for split in spec['replications']]
        for phase, split, default, mechanisms in contexts:
            seeds = PLAN['ancillary_seeds'] if phase == 'replication' else PLAN['main_seeds']
            result, predictions = collect(source, split, seeds, default=default, mechanisms=mechanisms)
            data = load_data(source, split)
            for name, record in result['models'].items():
                if name == 'neural_mean': continue  # exact API alias, not another method
                for seed, metrics in record['per_seed'].items():
                    pred = predictions[name][int(seed)]
                    macro = np.mean([pred['rr'][data.test[:,1] == r].mean()
                                     for r in np.unique(data.test[:,1])])
                    seed_rows.append([SOURCES[source], source, str(split), phase, name, NAMES.get(name,name),
                        int(seed), result['test_queries'], pred['metadata'].get('permutations',1),
                        *[metrics[key] for key in METRICS], float(macro)])
                for relation in record['relation_results']:
                    for seed, metrics in relation['per_seed'].items():
                        relation_rows.append([SOURCES[source], str(split), phase, NAMES.get(name,name),
                            int(seed), relation['relation'], relation['queries'], *[metrics[key] for key in METRICS]])
            if (source, split) not in seen_datasets:
                seen_datasets.add((source,split)); data=load_data(source, split); groups={}
                for i,(h,r,t) in enumerate(data.test):
                    group=groups.setdefault(ingredients(data,int(h)),len(groups))
                    query_rows.append([source,str(split),i,int(h),int(r),int(t),group,
                                       str(data.relations[int(r)]),len(data.range_candidates[int(r)])])
                obj={'audit':data.audit}
                metadata.extend([[f'dataset_{source}_{split}.json', *row] for row in flatten(obj)])
    phase_order={'main':0,'fixed':1,'replication':2}
    method_order={name:i for i,name in enumerate(('mean','neural_full_deepsets','neural_set_transformer','pairproto'))}
    seed_rows.sort(key=lambda row:(0 if row[1]=='PRIMARY' else 1,phase_order[row[3]],row[2],method_order.get(row[4],99),row[4],row[6]))
    relation_rows.sort(key=lambda row:(0 if row[0]=='TCM-MKG' else 1,phase_order[row[2]],row[1],row[3],row[5],row[4]))
    tables.append({'name':'SeedMetrics','headers':['Dataset','Source','Split','Evaluation','Model key','Method',
        'Seed','Queries','Permutations','MR','MRR','Hits@1','Hits@3','Hits@5','Hits@10','Relation macro MRR'], 'rows':seed_rows})
    summary=[]
    groups={}
    for i,row in enumerate(seed_rows,2): groups.setdefault(tuple(row[:6]),[]).append(i)
    for key, indexes in groups.items():
        assert indexes == list(range(min(indexes), max(indexes)+1))
        row=[key[0],key[2],key[3],key[5],len(indexes)]
        for col in ('J','K','L','M','N','O'):
            span=f"'SeedMetrics'!{col}{min(indexes)}:{col}{max(indexes)}"
            row += [f'=AVERAGE({span})',f'=STDEV.S({span})' if len(indexes)>1 else None]
        summary.append(row)
    tables.insert(0,{'name':'Summary','headers':['Dataset','Split','Evaluation','Method','Seeds']+
        [label for metric in ('MR','MRR','Hits@1','Hits@3','Hits@5','Hits@10') for label in (metric+' mean',metric+' SD')],
        'rows':summary,'formulas':True})
    tables.append({'name':'RelationMetrics','headers':['Dataset','Split','Evaluation','Method','Seed','Relation','Queries',
        'MR','MRR','Hits@1','Hits@3','Hits@5','Hits@10'],'rows':relation_rows})
    tables.append({'name':'Queries','headers':['Source','Split','Query index','Formula ID','Relation ID','Target ID',
        'Recipe group','Relation','Candidate count'],'rows':query_rows})
    recorded=[]
    excluded={'pairproto_paper_summary.json','pairproto_manuscript_static_audit.json',
              'pairproto_excel_validation.json','pairproto_excel_roundtrip.json'}
    for path in sorted(OUT.glob('*.json')):
        if path.name in excluded: continue
        obj=json.loads(path.read_text(encoding='utf-8'))
        metadata.extend([[path.name,*row] for row in flatten(obj)])
        source_rows.append([path.name,'json',hashlib.sha256(path.read_bytes()).hexdigest()])
        source,split,seed=identify(path.name)
        for pointer,metrics in recorded_metrics(obj):
            numbers=[]
            for key in METRICS:
                value=metrics.get(key)
                if value is None and key.startswith('hits'):value=metrics.get('hits@'+key[4:])
                numbers.append(value if isinstance(value,(int,float)) else None)
            recorded.append([path.name,pointer,source,split,seed,*numbers])
    schema={'schema_version':1,'report_date':'2026-10-23','metrics':list(METRICS),
            'protocol_sha256':hashlib.sha256((HERE/'pairproto_paper_protocol.json').read_bytes()).hexdigest(),
            'five_seed_amendment_sha256':hashlib.sha256((HERE/'pairproto_five_seed_amendment.json').read_bytes()).hexdigest() if (HERE/'pairproto_five_seed_amendment.json').exists() else None,
            'rank_convention':'Filtered average-tie rank; smaller MR is better; larger MRR/Hits is better',
            'source_of_truth':'Rank cells + Queries + typed Metadata; metric sheets are derived views',
            'recipe_scope':'Current paper main/fixed/replication; all available saved rank files also retained'}
    metadata.extend([['excel_schema.json',*row] for row in flatten(schema)])
    arrays=[]; vectors={}; groups={}; dedup={}; array_metrics=[]
    excluded_keys={'heads','relations','coverage','lengths','nearest_jaccard','target_recipe_support','recipe_groups','formula_ids'}
    for path in sorted(OUT.glob('*.npz')):
        source_rows.append([path.name,'npz',hashlib.sha256(path.read_bytes()).hexdigest()])
        with np.load(path) as z:
            for key in z.files:
                original=z[key];values=original.reshape(-1).copy();encoding='value'
                if key not in excluded_keys and values.dtype.kind=='f' and len(values) and np.all((values>0)&(values<=1)):
                    ranks=np.rint(2/values)/2
                    if np.allclose(1/values,ranks,rtol=0,atol=1e-7):
                        encoding='rank';values=ranks
                        dims=original.shape
                        for repeat,rr in enumerate(original.reshape((-1,dims[-1]))):
                            source,split,seed=identify(path.name)
                            array_metrics.append([path.name,key,source,split,seed,repeat,len(rr),*metric_values(rr)])
                signature=(str(values.dtype),len(values),hashlib.sha256(values.tobytes()).hexdigest())
                if signature not in dedup:
                    n=len(values);sheet=f'Ranks_{n}';group=groups.setdefault(n,[])
                    col=len(group)+2;ident=f'A{len(vectors)+1:04d}'
                    group.append(ident);vectors[ident]=values.tolist();dedup[signature]=(sheet,col,ident)
                sheet,col,ident=dedup[signature]
                arrays.append([path.name,key,sheet,col,str(original.dtype),json.dumps(original.shape),encoding,ident])
    tables.append({'name':'ArrayMetrics','headers':['Source file','Array','Source','Split','Seed','Repeat index','Queries',
        'MR','MRR','Hits@1','Hits@3','Hits@5','Hits@10'],'rows':array_metrics})
    tables.append({'name':'RecordedMetrics','headers':['Source file','Record path','Source','Split','Seed',
        'MR','MRR','Hits@1','Hits@3','Hits@5','Hits@10'],'rows':recorded})
    tables.append({'name':'ArrayIndex','headers':['Source file','Array','Worksheet','Column number','Dtype','Shape',
        'Cell encoding','Column ID'],'rows':arrays})
    tables.append({'name':'Metadata','headers':['Source file','JSON pointer','Type','Value'],'rows':metadata})
    tables.append({'name':'SourceFiles','headers':['Source file','Original format','Original SHA256'],'rows':source_rows})
    for n,identifiers in sorted(groups.items()):
        # Query positions are local to each source array; ArrayIndex identifies its
        # heads/relations and shape. Unrelated files are never joined by row alone.
        tables.append({'name':f'Ranks_{n}','headers':['Array position']+identifiers,
                       'rows':[[i]+[vectors[ident][i] for ident in identifiers] for i in range(n)]})
    tables.append({'name':'ReadMe','headers':['Topic','说明 / Definition'],'rows':[
        ['范围','Summary/SeedMetrics 包含主实验、固定配置及额外划分的全部种子；ArrayMetrics 另保留所有可用排名缓存和每次置换。'],
        ['权威输入','Ranks_*、Queries、ArrayIndex、Metadata 是论文计算输入；Summary 和 SeedMetrics 是方便阅读的派生结果。'],
        ['MR','Mean rank = mean(rank)，越低越好。'],['MRR','Mean reciprocal rank = mean(1/rank)，越高越好。'],
        ['Hits@k','mean(rank <= k)，k = 1, 3, 5, 10。'],
        ['并列','保留原评估的平均并列排名，可能出现半整数；不重新定义为悲观或乐观排名。'],
        ['置换','SeedMetrics 先在同一个训练种子内平均五次置换，再跨训练种子汇总；ArrayMetrics 保留每次置换。'],
        ['SD','Summary 使用训练种子之间的样本标准差；它不等于配方聚类置信区间。'],
        ['缺失','仅有汇总值、没有逐查询排名的历史/开发记录保留在 Metadata，不能凭空补算 MR。'],
        ['数值精度','排名存为精确整数/半整数，避免 Excel 小数舍入使 Hits@3 的边界发生变化。'],
        ['重算','运行 python build_pairproto_paper_artifacts.py --require-complete，从本工作簿重算论文表格、区间和结果图。'],
        ['新增实验','显式运行 python export_pairproto_excel.py 更新此工作簿；重算论文不会自行覆盖 Excel 或重新训练。'],
        ['编辑规则','更正输入应修改对应的排名/元数据单元格；只改派生指标表不会改变论文结果。'],
        ['限制','本表不包含预测概率，不能从排名推导 AUROC、AUPRC 或校准误差。'],
    ]})
    dest.mkdir(parents=True,exist_ok=True)
    payload={'tables':tables,'output':str(dest/'PairProto_All_Seeds.xlsx'),'preview_dir':str(dest/'previews')}
    path=dest/'excel_payload.json';path.write_text(json.dumps(payload,ensure_ascii=False,allow_nan=False),encoding='utf-8')
    print('EXCEL INPUT',len(seed_rows),'seed rows',len(array_metrics),'array/repeat rows',
          len(arrays),'arrays',sum(len(t['rows'])*len(t['headers']) for t in tables),'cells',flush=True)
    return path


def portable_writer(payload_path):
    """Portable fallback only where the bundled Artifact Tool is unavailable.

    openpyxl does not evaluate formulas; Excel evaluates the Summary on opening.
    Paper calculations read source ranks and do not depend on formula caches.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    payload=json.loads(payload_path.read_text(encoding='utf-8'))
    book=Workbook();book.remove(book.active)
    for spec in payload['tables']:
        sheet=book.create_sheet(spec['name']);sheet.append(spec['headers'])
        for row in spec['rows']:sheet.append(row)
        sheet.freeze_panes='C2' if not spec['name'].startswith('Ranks_') else 'B2'
        sheet.sheet_view.showGridLines=False
        for cell in sheet[1]:
            cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='243E56');cell.alignment=Alignment(wrap_text=True)
        sheet.row_dimensions[1].height=32
        for col in range(1,len(spec['headers'])+1):sheet.column_dimensions[get_column_letter(col)].width=18
        if spec['name'] in ('Summary','SeedMetrics','RelationMetrics'):
            sheet.column_dimensions['D' if spec['name']!='SeedMetrics' else 'F'].width=43
        if spec['name'] in ('ArrayMetrics','ArrayIndex','Metadata','SourceFiles','RecordedMetrics'):
            sheet.column_dimensions['A'].width=95
        if spec['name']=='Metadata':sheet.column_dimensions['B'].width=72;sheet.column_dimensions['D'].width=110
        if spec['name']=='ReadMe':sheet.column_dimensions['B'].width=120
        if spec['rows'] and not spec['name'].startswith('Ranks_'):
            table=Table(displayName='Evidence_'+spec['name'],ref=sheet.dimensions)
            table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            sheet.add_table(table)
    book.save(payload['output'])
    cache_portable_summary_values(payload_path)
    print('EXCEL SAVED (portable fallback; Summary caches verified from SeedMetrics; formulas remain editable)',payload['output'],flush=True)



def cache_portable_summary_values(payload_path):
    """Cache only declared Summary AVERAGE/STDEV.S formulas from SeedMetrics."""
    import re
    import zipfile
    import xml.etree.ElementTree as ET
    import numpy as np
    from openpyxl.utils import column_index_from_string
    payload=json.loads(payload_path.read_text(encoding='utf-8'))
    seed_rows=next(t['rows'] for t in payload['tables'] if t['name']=='SeedMetrics')
    workbook=Path(payload['output']);temporary=workbook.with_suffix('.cached.tmp')
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
        'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    with zipfile.ZipFile(workbook) as source:
        index=ET.fromstring(source.read('xl/workbook.xml'))
        summary=next(s for s in index.findall('m:sheets/m:sheet',ns) if s.get('name')=='Summary')
        identity=summary.get('{'+ns['r']+'}id')
        rels=ET.fromstring(source.read('xl/_rels/workbook.xml.rels'))
        target=next(x.get('Target') for x in rels if x.get('Id')==identity)
        member=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        xml=ET.fromstring(source.read(member));count=0
        pattern=re.compile(r"^(AVERAGE|STDEV\.S)\('SeedMetrics'!([A-Z]+)([0-9]+):\2([0-9]+)\)$")
        for cell in xml.findall('.//m:sheetData/m:row/m:c',ns):
            formula=cell.find('m:f',ns)
            if formula is None:continue
            match=pattern.fullmatch(formula.text or '')
            if not match:raise ValueError('Unsupported portable formula: '+str(formula.text))
            operation,column,first,last=match.groups()
            position=column_index_from_string(column)-1
            values=np.array([seed_rows[i-2][position] for i in range(int(first),int(last)+1)],dtype=float)
            summary_rows=next(t['rows'] for t in payload['tables'] if t['name']=='Summary')
            row_number=int(re.search(r'[0-9]+$',cell.get('r')).group())
            expected_count=int(summary_rows[row_number-2][4])
            assert len(values)==expected_count and expected_count>=2 and np.isfinite(values).all()
            number=float(values.mean()) if operation=='AVERAGE' else float(values.std(ddof=1))
            value=cell.find('m:v',ns)
            if value is None:value=ET.SubElement(cell,'{'+ns['m']+'}v')
            value.text=repr(number);cell.attrib.pop('t',None);count+=1
        replacement=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(temporary,'w') as result:
            for item in source.infolist():
                result.writestr(item,replacement if item.filename==member else source.read(item.filename))
    temporary.replace(workbook)
    print('PORTABLE SUMMARY FORMULA CACHES',count,flush=True)


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--output-dir',type=Path,default=DEFAULT_DEST)
    parser.add_argument('--payload-only',action='store_true');args=parser.parse_args()
    payload=make_payload(args.output_dir)
    if not args.payload_only:
        node=os.environ.get('PAIRPROTO_ARTIFACT_NODE')
        available=False
        if node and Path(node).is_file():
            check=subprocess.run([node,'--input-type=module','-e',"import '@oai/artifact-tool'"],cwd=HERE,
                                 stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
            available=check.returncode==0
        if available:
            subprocess.run([node,'--max-old-space-size=12000',str(HERE/'build_pairproto_excel.mjs'),str(payload)],check=True)
        else:
            print('Bundled Artifact Tool unavailable: using the documented portable writer.',flush=True)
            portable_writer(payload)


if __name__=='__main__': main()
