"""Optional public fallback when Artifact Tool is unavailable."""
from pathlib import Path
from collections import defaultdict
import json,zipfile,xml.etree.ElementTree as ET
def write(payload_path):
    import numpy as np
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    payload=json.loads(Path(payload_path).read_text(encoding="utf-8"))
    wb=Workbook();summary=wb.active;summary.title="SR_Results"
    for spec in payload["tables"]:
        sheet=wb.create_sheet(spec["name"]);sheet.append(spec["headers"])
        for row in spec["rows"]:sheet.append(row)
        sheet.sheet_view.showGridLines=False;sheet.freeze_panes="E2"
        for cell in sheet[1]:
            cell.font=Font(name="Arial",bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="263F56")
        for column in range(1,len(spec["headers"])+1):sheet.column_dimensions[get_column_letter(column)].width=20
    summary.append(["PairProto-KGE Scientific Reports evidence"])
    summary.append(["Source","Split","Evaluation","Model","n","MRR","SD MRR","Hits@1","SD H1","Hits@3","SD H3","Hits@10","SD H10","MR","SD MR","Hits@5","SD H5"])
    rows=next(t["rows"] for t in payload["tables"] if t["name"]=="SR_Seeds")
    groups=defaultdict(list)
    for index,row in enumerate(rows,2):groups[tuple(row[1:6])].append(index)
    cached={}
    for group,indices in sorted(groups.items()):
        source,split,phase,key,label=group
        values=[source,split,phase,label,len(indices)];r=summary.max_row+1
        for column in [11,12,13,15,10,14]:
            refs=",".join("'SR_Seeds'!"+get_column_letter(column)+str(i) for i in indices)
            numbers=np.asarray([rows[i-2][column-1] for i in indices],dtype=float)
            for operation,number in [("AVERAGE",float(numbers.mean())),("STDEV.S",float(numbers.std(ddof=1)))]:
                address=get_column_letter(len(values)+1)+str(r);cached[address]=number
                values.append("="+operation+"("+refs+")")
        summary.append(values)
    summary.freeze_panes="E3";summary.column_dimensions["D"].width=43
    wb.save(payload["output"])
    ns="{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    target=Path(payload["output"]);tmp=target.with_suffix(".cached.tmp")
    with zipfile.ZipFile(target) as source:
        xml=ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
        for cell in xml.iter(ns+"c"):
            if cell.get("r") in cached:
                v=cell.find(ns+"v")
                if v is None:v=ET.SubElement(cell,ns+"v")
                v.text=repr(cached[cell.get("r")])
        with zipfile.ZipFile(tmp,"w",zipfile.ZIP_DEFLATED) as result:
            for name in source.namelist():result.writestr(name,ET.tostring(xml,encoding="utf-8",xml_declaration=True) if name=="xl/worksheets/sheet1.xml" else source.read(name))
    tmp.replace(target)
    print("PORTABLE FOLLOWUP WORKBOOK EXPORTED",target,flush=True)
