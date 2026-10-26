"""Verify Excel round-trip precision and the paper's exclusive Excel dependency."""
from __future__ import annotations
import argparse
import json
from pathlib import Path
from unittest.mock import patch
import numpy as np

from pairproto_excel_source import evidence, RESULTS_DIR, workbook_path
from pairproto_evidence import build_report, collect, config_for, PLAN
from pair_prototype_full_suite import neural_tag


def compare_existing(old, new, where='report'):
    """Preserve every old result; allow newly exposed ranking metrics/provenance."""
    if isinstance(old,dict):
        for key,value in old.items():
            if key not in new: raise AssertionError(f'Missing result {where}/{key}')
            compare_existing(value,new[key],where+'/'+key)
    elif isinstance(old,list):
        if len(old)!=len(new): raise AssertionError(where+' length')
        for i,(a,b) in enumerate(zip(old,new)): compare_existing(a,b,where+'/'+str(i))
    elif isinstance(old,(int,float)) and not isinstance(old,bool):
        if not np.isclose(old,new,rtol=1e-11,atol=1e-12): raise AssertionError(f'{where}: {old} != {new}')
    elif old!=new: raise AssertionError(f'{where}: {old!r} != {new!r}')


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--baseline',type=Path)
    parser.add_argument('--compare-raw-arrays',action='store_true');args=parser.parse_args()
    store=evidence();checked=0
    if args.compare_raw_arrays:
        for filename,arrays in store.arrays.items():
            with np.load(RESULTS_DIR/filename) as original:
                assert set(original.files)==set(arrays),filename
                for key,value in arrays.items():
                    assert original[key].shape==value.shape,(filename,key)
                    assert np.allclose(original[key],value,rtol=1e-13,atol=1e-14),(filename,key)
                    checked+=1
    original_read=Path.read_text
    def guarded_read(path,*a,**kw):
        if path.resolve().parent==RESULTS_DIR.resolve():
            raise AssertionError('Paper attempted raw JSON read: '+str(path))
        return original_read(path,*a,**kw)
    with patch.object(Path,'read_text',guarded_read), \
         patch('pair_prototype_full_suite.load_data',side_effect=AssertionError('Raw dataset access')), \
         patch('numpy.load',side_effect=AssertionError('Raw NPZ access')):
        report=build_report()
        assert report['complete'],report['missing']
        # A controlled rank perturbation must change the recomputed metric.
        source='PRIMARY';split=PLAN['datasets'][source]['main'];seed=PLAN['main_seeds'][0]
        cfg=config_for(source,'full_deepsets')
        tag=neural_tag(source,split,'full_deepsets',seed,cfg['lr'],cfg['heads'])+'_rr.npz'
        rr=store.arrays[tag]['rr'];old=float(rr[0]);new=1. if old!=1. else .5
        before=report['sources'][source]['main']['models']['neural_full_deepsets']['per_seed'][str(seed)]['mrr']
        try:
            rr[0]=new
            changed,_=collect(source,split,[seed])
            after=changed['models']['neural_full_deepsets']['per_seed'][str(seed)]['mrr']
            assert np.isclose(after-before,(new-old)/len(rr),rtol=1e-9,atol=1e-14)
        finally:rr[0]=old
    if args.baseline:
        baseline=json.loads(args.baseline.read_text(encoding='utf-8'))
        # Export changes only the input provenance; all scientific results must match.
        assert report['input_provenance']['mode']=='excel'
        baseline_results={k:v for k,v in baseline.items() if k!='input_provenance'}
        report_results={k:v for k,v in report.items() if k!='input_provenance'}
        compare_existing(baseline_results,report_results)
    from openpyxl import load_workbook
    book=load_workbook(workbook_path(),read_only=True,data_only=True)
    n=0
    for row in book['SeedMetrics'].iter_rows(min_row=2,values_only=True):
        _,source,split,phase,key,_,seed,queries,permutations,*metrics=row
        resource=report['sources'][source]
        section=resource['replications'][str(split)] if phase=='replication' else resource[phase]
        actual=section['models'][key]['per_seed'][str(seed)]
        for metric,value in zip(('mr','mrr','hits1','hits3','hits5','hits10'),metrics):
            assert np.isclose(value,actual[metric],rtol=1e-11,atol=1e-12),(source,key,seed,metric)
        n+=1
    for row in book['Summary'].iter_rows(min_row=2,values_only=True):
        if report['protocol'].get('five_seed_amendment'):
            assert row[4]==5,('Non-five-seed paper comparison',row[:5])
        assert all(isinstance(value,(int,float)) for value in row[4:]),'Missing formula cache'
    book.close()
    result={'workbook':str(workbook_path()),'sha256':store.sha256,'rank_arrays_checked':checked,
            'seed_rows_verified':n,'complete':report['complete'],
            'no_raw_json_npz_or_dataset_reads':'pass','rank_perturbation_changes_metric':'pass',
            'pre_excel_results_preserved':'pass' if args.baseline else 'not requested',
            'cached_summary_formula_values':'pass'}
    (RESULTS_DIR/'pairproto_excel_validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2),flush=True)


if __name__=='__main__':main()
