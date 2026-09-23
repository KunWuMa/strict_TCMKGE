# -*- coding: utf-8 -*-
"""Inspect candidate TCM knowledge-graph source tables without modifying them."""

from __future__ import annotations

import csv
import json
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
PART1 = next(ROOT.glob("01_*"))


def decode(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError(path)


def inspect_tsv(path: Path) -> dict:
    text = decode(path)
    rows = list(csv.reader(text.splitlines(), delimiter="\t"))
    width = max((len(row) for row in rows), default=0)
    return {
        "file": path.name,
        "rows_including_header": len(rows),
        "width": width,
        "header": rows[0] if rows else [],
        "examples": rows[1:4],
    }


def inspect_syndrome_knowledge(path: Path) -> dict:
    rows = [json.loads(line) for line in decode(path).splitlines() if line.strip()]
    return {
        "file": path.name,
        "records": len(rows),
        "fields": list(rows[0]) if rows else [],
        "examples": rows[:3],
    }


def inspect_symmap(path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        return {"file": path.name, "error": f"openpyxl unavailable: {exc}"}
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    sample = []
    for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
        sample.append([value for value in row])
    return {
        "file": path.name,
        "sheet": sheet.title,
        "max_row": sheet.max_row,
        "max_column": sheet.max_column,
        "sample": sample,
    }


def main() -> None:
    result = {
        "tcm_mkg": [inspect_tsv(path) for path in sorted((PART1 / "TCM-MKG").glob("*.tsv"))],
        "syndrome_knowledge": inspect_syndrome_knowledge(PART1 / "TCM-SD" / "syndrome_knowledge.json"),
        "symmap": [inspect_symmap(path) for path in sorted((PART1 / "SymMap_v2").glob("* file.xlsx"))],
    }
    out = Path(__file__).with_name("source_audit.json")
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "tcm_mkg_files": len(result["tcm_mkg"]),
        "tcm_mkg_rows": sum(max(0, item["rows_including_header"] - 1) for item in result["tcm_mkg"]),
        "syndrome_knowledge_records": result["syndrome_knowledge"]["records"],
        "symmap_files": len(result["symmap"]),
        "output": str(out),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
